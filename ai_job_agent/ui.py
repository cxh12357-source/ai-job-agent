"""Streamlit home for the end-to-end Charles-ai-job-agent MVP.

The page deliberately keeps the happy path linear: resume -> jobs -> queue ->
review -> explicit confirmation.  Existing advanced Greenhouse tools remain in
separate tabs in ``app.py``.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Mapping
from pathlib import Path
from typing import Any

import streamlit as st

from branding import APP_NAME
from config import SETTINGS
from job_assistant.autofill import AutofillError, create_autofill_plan
from job_assistant.profile import ProfileError, load_profile as load_saved_applicant_profile
from job_assistant.profile import select_resume as select_saved_resume

from .browser.agent import BrowserAgent, BrowserRunResult
from .browser.demo_server import DemoServer
from .company_scope import company_scope_for_job, filter_jobs_by_scope, scope_label
from .database import DuplicateApplicationError
from .demo_data import demo_profile
from .models import CandidateProfile, JobPosting, MatchAssessment
from .services.application_service import ApplicationService
from .services.autofill_monitor import (
    AutofillLaunchRequest,
    is_autofill_scheduled,
    schedule_autofill_batch,
)
from .services.early_career import filter_early_career_jobs
from .services.job_matcher import match_job
from .services.job_searcher import (
    DEFAULT_GREENHOUSE_TOKENS,
    JobSearchError,
    search_jobs_detailed,
)
from .services.local_job_import import (
    LocalJobImportError,
    local_job_export_template,
    parse_local_job_export,
)
from .services.local_scraper_bundle import build_local_scraper_bundle
from .services.official_source_searcher import (
    OfficialCareerSearchError,
    search_official_career_urls,
)
from .services.profile_builder import build_candidate_profile
from .services.resume_intake import process_resume_upload
from .services.resume_parser import ResumeParseError, parse_resume
from .sources.campus_catalog import (
    CAMPUS_CAREER_SITES,
    DEFAULT_CAMPUS_LABELS,
    MAX_CAMPUS_SOURCES_PER_SEARCH,
    campus_sites_for_regions,
    campus_sites_for_scope,
    resolve_campus_selection,
)
from .sources.catalog import CAREER_SITE_BY_LABEL
from .ui_state import clear_job_search_state, clear_resume_state
from .ui_text import markdown_literal


STATUS_LABELS = {
    "pending": "等待处理",
    "opening": "正在打开官网",
    "filling": "正在填写",
    "waiting_user": "待你在官网补充",
    "ready_to_submit": "准备提交",
    "submitted": "已投递",
    "failed": "失败",
    "skipped": "已跳过",
}


@st.cache_resource(show_spinner=False)
def _application_service() -> ApplicationService:
    SETTINGS.ensure_local_directories()
    return ApplicationService(
        SETTINGS.database_path,
        demo_mode=SETTINGS.demo_mode,
        auto_submit=False,
    )


def _init_state() -> None:
    defaults: dict[str, Any] = {
        "aja_profile": None,
        "aja_resume_text": "",
        "aja_resume_path": "",
        "aja_resume_filename": "",
        "aja_resume_size": 0,
        "aja_resume_upload_revision": 0,
        "aja_jobs": [],
        "aja_assessments": {},
        "aja_selected": set(),
        "aja_search_report": {},
        "aja_notice": "",
        "aja_legacy_queue_migrated": False,
        "aja_saved_profile_checked": False,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def _clear_job_search_state() -> None:
    clear_job_search_state(st.session_state)


def _forget_resume(*, reset_upload: bool = False) -> None:
    clear_resume_state(st.session_state, reset_upload=reset_upload)


def _inject_style() -> None:
    st.markdown(
        """
        <style>
        .block-container {max-width: 1180px; padding-top: 1.4rem; padding-bottom: 4rem;}
        [data-testid="stMetric"] {
            background: rgba(247,248,250,.82); border: 1px solid #e8eaed;
            border-radius: 14px; padding: .8rem 1rem;
        }
        .aja-hero {
            padding: 1.6rem 1.7rem; border: 1px solid #e8eaed; border-radius: 20px;
            background: linear-gradient(135deg, #ffffff 12%, #f4f7ff 100%);
            margin-bottom: 1rem;
        }
        .aja-eyebrow {font-size: .78rem; letter-spacing: .12em; color: #667085; font-weight: 700;}
        .aja-hero h1 {font-size: 2.25rem; margin: .25rem 0 .35rem 0; letter-spacing: -.04em;}
        .aja-hero p {max-width: 760px; color: #5f6672; margin: 0; line-height: 1.65;}
        .aja-flow {margin-top: .9rem; color: #334155; font-size: .9rem;}
        .aja-safe {color: #16794b; font-weight: 650;}
        .aja-score {font-size: 1.8rem; font-weight: 760; letter-spacing: -.04em;}
        .aja-muted {color: #667085; font-size: .88rem;}
        .aja-chip {display:inline-block; padding:.16rem .5rem; margin:.12rem .18rem .12rem 0;
            border:1px solid #e5e7eb; border-radius:999px; font-size:.78rem; background:#fafafa;}
        </style>
        """,
        unsafe_allow_html=True,
    )


def _profile_from_state() -> CandidateProfile | None:
    value = st.session_state.aja_profile
    if isinstance(value, CandidateProfile):
        return value
    if isinstance(value, dict):
        return CandidateProfile.model_validate(value)
    return None


def _saved_profile_name(first_name: str, last_name: str) -> str | None:
    first = str(first_name or "").strip()
    last = str(last_name or "").strip()
    if not first and not last:
        return None
    if all("\u4e00" <= char <= "\u9fff" for char in first + last):
        return f"{last}{first}"
    return " ".join(part for part in (first, last) if part)


def _restore_saved_candidate_profile() -> bool:
    """Restore a previously approved local profile without calling an external API."""

    if SETTINGS.demo_mode or st.session_state.aja_saved_profile_checked:
        return False
    st.session_state.aja_saved_profile_checked = True
    try:
        applicant = load_saved_applicant_profile()
        try:
            resume_path = select_saved_resume(applicant, "zh")
        except ProfileError:
            resume_path = select_saved_resume(applicant, "en", company_foreign=True)
        content = resume_path.read_bytes()
        resume_text = parse_resume(resume_path.name, content)
        parsed = build_candidate_profile(resume_text, use_openai=False)
        warnings = list(parsed.parse_warnings)
        if applicant.email and parsed.email and applicant.email.casefold() != parsed.email.casefold():
            warnings.append("申请资料中的邮箱与简历不一致；官网预填使用你在“申请资料”中保存的邮箱")
        profile_phone = "".join(character for character in str(applicant.phone or "") if character.isdigit())
        resume_phone = "".join(character for character in str(parsed.phone or "") if character.isdigit())
        if profile_phone and resume_phone and profile_phone != resume_phone:
            warnings.append("申请资料中的手机号与简历不一致；官网预填使用你在“申请资料”中保存的手机号")
        restored = parsed.model_copy(
            update={
                "name": parsed.name
                or _saved_profile_name(applicant.first_name, applicant.last_name),
                "email": applicant.email or parsed.email or None,
                "phone": applicant.phone or parsed.phone or None,
                "location": applicant.location or parsed.location or None,
                "parse_warnings": warnings,
            }
        )
    except (OSError, ProfileError, ResumeParseError, ValueError):
        return False
    if not all((restored.name, restored.email, restored.phone)):
        return False
    st.session_state.aja_profile = restored
    st.session_state.aja_resume_text = resume_text
    st.session_state.aja_resume_path = str(resume_path)
    st.session_state.aja_notice = "已自动载入你保存在本机的申请资料和简历。"
    return True


def _job_key(job: JobPosting) -> str:
    return str(job.job_id or job.job_url)


def _name_parts(name: str | None) -> tuple[str, str]:
    value = (name or "").strip()
    if not value:
        return "", ""
    pieces = value.split()
    if len(pieces) > 1:
        return " ".join(pieces[:-1]), pieces[-1]
    if len(value) in {2, 3, 4} and all("\u4e00" <= char <= "\u9fff" for char in value):
        return value[1:], value[0]
    return value, value


def _facts_line(*values: object) -> str:
    return "｜".join(str(value).strip() for value in values if str(value or "").strip())


def _experience_summary(items: object) -> str:
    result: list[str] = []
    for item in items or ():
        header = _facts_line(
            getattr(item, "company", None),
            getattr(item, "title", None),
            getattr(item, "location", None),
            _facts_line(getattr(item, "start_date", None), getattr(item, "end_date", None)),
        )
        details = [str(getattr(item, "description", None) or "").strip()]
        details.extend(str(value).strip() for value in (getattr(item, "achievements", None) or ()))
        details = [value for value in details if value]
        block = "\n".join([header, *details]).strip()
        if block:
            result.append(block)
    return "\n\n".join(result)


def _project_summary(items: object) -> str:
    result: list[str] = []
    for item in items or ():
        header = _facts_line(
            getattr(item, "name", None),
            getattr(item, "role", None),
            _facts_line(getattr(item, "start_date", None), getattr(item, "end_date", None)),
        )
        details = [str(getattr(item, "description", None) or "").strip()]
        details.extend(str(value).strip() for value in (getattr(item, "achievements", None) or ()))
        technologies = _facts_line(*(getattr(item, "technologies", None) or ()))
        if technologies:
            details.append(technologies)
        block = "\n".join([header, *(value for value in details if value)]).strip()
        if block:
            result.append(block)
    return "\n\n".join(result)


def _education_summary(profile: CandidateProfile) -> str:
    result: list[str] = []
    entries = list(profile.education or ())
    if not entries and any((profile.school, profile.degree, profile.major, profile.graduation_date)):
        line = _facts_line(
            profile.school,
            profile.degree,
            profile.major,
            profile.graduation_date,
        )
        return line
    for index, item in enumerate(entries):
        line = _facts_line(
            profile.school if index == 0 else item.school,
            profile.degree if index == 0 else item.degree,
            profile.major if index == 0 else item.major,
            item.location,
            _facts_line(
                item.start_date,
                profile.graduation_date if index == 0 else item.graduation_date,
            ),
            item.description,
        )
        if line:
            result.append(line)
    return "\n".join(result)


def _profile_values(profile: CandidateProfile) -> dict[str, str]:
    first_name, last_name = _name_parts(profile.name)
    language_summary = "\n".join(
        _facts_line(item.name, item.proficiency) for item in (profile.languages or ())
    )
    certification_summary = "\n".join(
        _facts_line(item.name, item.issuer, item.issue_date, item.expiry_date)
        for item in (profile.certifications or ())
    )
    award_summary = "\n".join(
        _facts_line(item.name, item.issuer, item.date, item.description)
        for item in (profile.awards or ())
    )
    return {
        "first_name": first_name,
        "last_name": last_name,
        "full_name": profile.name or "",
        "email": profile.email or "",
        "phone": profile.phone or "",
        "location": profile.location or "",
        "school": profile.school or "",
        "degree": profile.degree or "",
        "major": profile.major or "",
        "graduation_date": profile.graduation_date or "",
        "self_introduction": profile.self_introduction or "",
        "education": _education_summary(profile),
        "skills": "、".join(profile.skills or ()),
        "languages": language_summary,
        "certifications": certification_summary,
        "internships": _experience_summary(profile.internships),
        "work_experience": _experience_summary(profile.work_experience),
        "projects": _project_summary(profile.projects),
        "awards": award_summary,
        "resume": str(st.session_state.aja_resume_path or ""),
    }


def _ensure_demo_resume() -> Path:
    """Create a harmless local DOCX used only by the loopback Demo form."""

    target = SETTINGS.uploads_dir / "demo_candidate_resume.docx"
    if target.exists():
        return target
    SETTINGS.uploads_dir.mkdir(parents=True, exist_ok=True)
    from docx import Document

    document = Document()
    document.add_heading("Demo Candidate", level=1)
    document.add_paragraph(f"Local {APP_NAME} demo resume. Not a real candidate.")
    document.add_paragraph("Skills: Python, Data Analysis, LLM, AI Agent")
    document.save(target)
    return target


def _run_local_demo_browser(
    values: dict[str, str], *, confirmed: bool
) -> BrowserRunResult:
    """Exercise the real Playwright stack against the loopback-only form."""

    with DemoServer() as server, BrowserAgent(
        headless=True,
        allow_local_demo=True,
        screenshot_dir=SETTINGS.logs_dir,
    ) as agent:
        return agent.run_application(
            server.apply_url,
            values,
            environment="demo",
            review_confirmed=confirmed,
            # This never submits externally.  The browser layer only permits its
            # hard-coded local demo control, and only after explicit confirmation.
            auto_submit=confirmed,
        )


def _start_visible_official_application(
    service: ApplicationService,
    queue_item: dict[str, Any],
    profile: CandidateProfile,
) -> bool:
    """Open one selected official application and safely prefill known facts."""

    request = _build_official_autofill_request(service, queue_item, profile)
    handle = schedule_autofill_batch(service, [request])
    return bool(handle.scheduled_queue_ids)


def _build_official_autofill_request(
    service: ApplicationService,
    queue_item: Mapping[str, Any],
    profile: CandidateProfile,
) -> AutofillLaunchRequest:
    """Validate one queue item and build its value-private browser request."""

    if not SETTINGS.demo_mode and profile.parse_method == "demo":
        raise AutofillError("正式投递必须先上传并解析你的真实简历")
    queue_id = int(queue_item["id"])
    raw_job = queue_item.get("job")
    job = dict(raw_job) if isinstance(raw_job, Mapping) else None
    if job is None:
        application = queue_item.get("application")
        application_job_pk = (
            application.get("job_pk") if isinstance(application, Mapping) else None
        )
        job_pk = queue_item.get("job_pk") or application_job_pk
        job = service.database.get_job(int(job_pk)) if job_pk is not None else None
    if job is None:
        raise AutofillError("岗位记录已失效，请重新搜索")
    resume_path = str(st.session_state.aja_resume_path or "").strip()
    plan = create_autofill_plan(
        str(job["job_url"]),
        _profile_values(profile),
        resume_path,
        source=str(job.get("source") or ""),
        consent=True,
        dry_run=False,
    )
    return AutofillLaunchRequest(queue_id=queue_id, plan=plan)


def _dashboard(service: ApplicationService) -> None:
    counts = service.dashboard_counts(recommended_threshold=70)
    columns = st.columns(6)
    values = (
        ("岗位", counts.get("total_jobs", 0)),
        ("推荐", counts.get("recommended_jobs", counts.get("recommended", 0))),
        ("待投递", counts.get("pending", counts.get("active_pending", 0))),
        ("已投递", counts.get("submitted", 0)),
        ("失败", counts.get("failed", 0)),
        ("待你补充", counts.get("needs_user_confirmation", 0)),
    )
    for column, (label, value) in zip(columns, values, strict=True):
        column.metric(label, int(value or 0))


def _resume_step() -> None:
    st.markdown("### 1 · 上传个人 PDF 简历")
    left, right = st.columns([1.55, 1], gap="large")
    with left:
        st.caption(
            "支持 PDF / DOCX，单个文件不超过 10 MB。"
            + ("文件会上传至 Streamlit 云端；本入口不将原文件写入云端磁盘。" if SETTINGS.cloud_deployment
               else "原文件保存在本机 uploads/，用于后续经你确认的官网预填。")
        )
        upload = st.file_uploader(
            "上传个人 PDF 简历",
            type=["pdf", "docx"],
            key=f"aja_resume_upload_{st.session_state.aja_resume_upload_revision}",
            on_change=_forget_resume,
            help="兼容 DOCX；扫描版 PDF 暂不包含 OCR。加密 PDF 请先在本机解密。",
        )
        use_ai = st.checkbox(
            "允许将简历文字发送给 OpenAI，辅助结构化解析",
            value=False,
            disabled=not SETTINGS.openai_enabled,
            key="aja_resume_ai_consent",
            help="默认关闭；关闭时只在运行 App 的服务器上解析，不调用外部 AI。",
        )
        action_a, action_b = st.columns(2)
        parse_clicked = action_a.button(
            "解析我的简历", type="primary", use_container_width=True,
            disabled=upload is None,
        )
        demo_clicked = action_b.button(
            "使用示例简历" if SETTINGS.demo_mode else "示例简历（仅 Demo）",
            use_container_width=True,
            disabled=not SETTINGS.demo_mode,
        )

        if parse_clicked and upload is not None:
            try:
                content = upload.getvalue()
                with st.spinner("正在提取事实并生成候选人档案…"):
                    result = process_resume_upload(
                        upload.name,
                        content,
                        cloud_mode=SETTINGS.cloud_deployment,
                        uploads_dir=SETTINGS.uploads_dir,
                        use_openai=bool(use_ai and SETTINGS.openai_enabled),
                        model=SETTINGS.openai_model,
                    )
                _forget_resume()
                st.session_state.aja_resume_text = result.text
                st.session_state.aja_resume_path = result.saved_path
                st.session_state.aja_resume_filename = result.filename
                st.session_state.aja_resume_size = result.size_bytes
                st.session_state.aja_profile = result.profile
                st.rerun()
            except (ResumeParseError, ValueError, OSError) as exc:
                _forget_resume()
                st.error(str(exc))

        if demo_clicked:
            _forget_resume(reset_upload=True)
            st.session_state.aja_profile = demo_profile()
            st.session_state.aja_resume_text = "Demo profile generated from local sample facts."
            st.session_state.aja_resume_path = "" if SETTINGS.cloud_deployment else str(_ensure_demo_resume())
            st.session_state.aja_resume_filename = "demo_candidate_resume.docx"
            st.rerun()

        current_profile = _profile_from_state()
        if current_profile is not None:
            filename = st.session_state.aja_resume_filename or "已载入的简历"
            st.success(f"已解析：{filename} · {len(st.session_state.aja_resume_text):,} 字符")
            with st.expander("查看 PDF / DOCX 提取文字"):
                st.text(st.session_state.aja_resume_text)
            st.button(
                "清除此会话简历",
                on_click=_forget_resume,
                kwargs={"reset_upload": True},
                help="清空本会话的档案、提取文字和匹配结果；不删除已有投递记录或本机历史文件。",
            )

    with right:
        st.info(
            "**安全边界**\n\n"
            "只提取简历中存在的事实；GPA、薪资、工作许可等未知信息不会猜。"
            "验证码、登录和网站限制会暂停。最终提交前始终需要你确认。"
        )
        mode = "Demo Mode · 不会向外部网站提交" if SETTINGS.demo_mode else "Real Mode · 提交前人工确认"
        st.caption(mode)
        if SETTINGS.cloud_deployment:
            st.caption("提取文字和候选人档案仅供当前会话使用；请自行保存原 PDF。云端不代你向招聘网站上传。")


def _profile_step(profile: CandidateProfile) -> CandidateProfile:
    st.markdown("### 2 · 核对目标")
    with st.container(border=True):
        top = st.columns([1.2, 1, 1])
        top[0].write(f"**{profile.name or '姓名待确认'}**")
        top[0].caption(profile.email or "邮箱待确认")
        top[1].write(profile.school or "学校待确认")
        top[1].caption(" / ".join(filter(None, [profile.degree, profile.major])) or "学历与专业待确认")
        top[2].write(f"已识别 {len(profile.skills or [])} 项技能")
        top[2].caption("AI 解析" if profile.parse_method == "openai" else "本地安全解析")

        if profile.parse_warnings:
            st.warning("；".join(profile.parse_warnings))

        with st.expander("核对官网自动填写资料", expanded=False):
            st.caption("这些值会在你点击一键投递后写入对应官网字段；空白项和敏感问题不会猜。")
            p1, p2, p3 = st.columns(3)
            applicant_name = p1.text_input("姓名", value=profile.name or "")
            applicant_email = p2.text_input("邮箱", value=profile.email or "")
            applicant_phone = p3.text_input("手机号", value=profile.phone or "")
            e1, e2, e3, e4 = st.columns([1.3, 0.8, 1.15, 0.9])
            applicant_school = e1.text_input("学校", value=profile.school or "")
            applicant_degree = e2.text_input("学历", value=profile.degree or "")
            applicant_major = e3.text_input("专业", value=profile.major or "")
            applicant_graduation = e4.text_input(
                "毕业时间", value=profile.graduation_date or "", placeholder="YYYY-MM"
            )
            applicant_summary = st.text_area(
                "个人概述",
                value=profile.self_introduction or "",
                height=90,
                help="只填写简历已有内容；岗位动机等开放题不会直接套用这里。",
            )

        role_default = "，".join(profile.target_roles or [])
        location_default = "，".join(profile.target_locations or [])
        c1, c2, c3 = st.columns([1.4, 1.2, 1])
        roles = c1.text_input("目标岗位", value=role_default, placeholder="AI 应用，数据分析")
        locations = c2.text_input("就业地点", value=location_default, placeholder="上海，苏州")
        salary = c3.text_input(
            "目标薪资（可留空）", value=profile.expected_salary or "", placeholder="例如 15k-20k"
        )
        stage_labels = {
            "应届生 / 校招": "fresh_graduate",
            "在校生 / 实习": "internship",
            "有工作经验 / 社招": "experienced",
        }
        stage_by_value = {value: label for label, value in stage_labels.items()}
        d1, d2 = st.columns([1, 2.2])
        stage_label = d1.selectbox(
            "求职身份",
            tuple(stage_labels),
            index=tuple(stage_labels).index(
                stage_by_value.get(profile.career_stage, "应届生 / 校招")
            ),
            help="已按你的需求默认选择应届生；不会自动猜测具体毕业届次。",
        )
        current_location = d2.text_input(
            "当前所在地（官网表单用，可留空）",
            value=profile.location or "",
            placeholder="例如 上海",
            help="只在你明确填写时用于官网申请表，不会根据目标就业城市自动猜测。",
        )
        updated = profile.model_copy(
            update={
                "target_roles": [item.strip() for item in roles.replace("，", ",").split(",") if item.strip()] or None,
                "target_locations": [item.strip() for item in locations.replace("，", ",").split(",") if item.strip()] or None,
                "career_stage": stage_labels[stage_label],
                "expected_salary": salary.strip() or None,
                "location": current_location.strip() or None,
                "name": applicant_name.strip() or None,
                "email": applicant_email.strip() or None,
                "phone": applicant_phone.strip() or None,
                "school": applicant_school.strip() or None,
                "degree": applicant_degree.strip() or None,
                "major": applicant_major.strip() or None,
                "graduation_date": applicant_graduation.strip() or None,
                "self_introduction": applicant_summary.strip() or None,
            }
        )
        st.session_state.aja_profile = updated
        st.download_button(
            "下载解析档案 JSON",
            data=updated.model_dump_json(indent=2),
            file_name="candidate_profile.json",
            mime="application/json",
            help="下载刚刚核对后的最新档案。包含个人信息，请只保存到可信设备；不要把它当作岗位文件上传。",
        )
        with st.expander("查看解析详情与待确认项"):
            st.json(updated.model_dump(mode="json"), expanded=False)
            if updated.parse_warnings:
                st.warning("；".join(updated.parse_warnings))
    return updated


def _local_jobs_step(profile: CandidateProfile, service: ApplicationService) -> None:
    st.info(
        "本机运行 Playwright → 导出岗位 JSON → 在这里上传并匹配。"
        "网页不会连接或远程控制你的电脑，也不接收浏览器 Cookie。"
    )
    with st.expander("下载本地抓取脚本与使用说明", expanded=True):
        st.caption("仅供已获准自动读取的招聘页面使用。登录和验证码由你手动完成；网站禁止或限流时停止。")
        try:
            st.download_button(
                "下载本地 Playwright 工具包",
                data=build_local_scraper_bundle(),
                file_name="charles-local-scraper.zip",
                mime="application/zip",
                key="aja_download_local_scraper",
            )
        except OSError:
            st.warning("本地工具包暂不可用；可以从项目仓库获取 local_scraper.py。")
        st.code(
            "python -m pip install -r requirements-local-scraper.txt\n"
            "python -m playwright install chromium\n"
            "python local_scraper.py --demo --output output/local_jobs.json",
            language="bash",
        )
        st.caption("先运行 Demo 验证浏览器。真实网站请按工具包说明修改配置并明确确认读取权限。")
        st.download_button(
            "下载岗位 JSON 格式样例（非真实岗位）",
            data=local_job_export_template(),
            file_name="local_jobs.example.json",
            mime="application/json",
        )
    upload = st.file_uploader(
        "上传本地抓取的岗位 JSON",
        type=["json"],
        key="aja_local_jobs_upload",
        on_change=_clear_job_search_state,
        help="只上传 local_scraper.py 导出的岗位文件，最多 5 MB / 100 个岗位。不要上传简历档案、Cookie 或密码文件。",
    )
    if st.button("导入岗位并匹配", type="primary", disabled=upload is None):
        try:
            with st.spinner("正在校验岗位文件并计算匹配度…"):
                result = parse_local_job_export(upload.name, upload.getvalue())
                assessments = {}
                for job in result.jobs:
                    assessment = match_job(profile, job, use_openai=False)
                    assessments[_job_key(job)] = assessment
                    service.register_job(job, assessment)
            st.session_state.aja_jobs = result.jobs
            st.session_state.aja_assessments = assessments
            st.session_state.aja_selected = set()
            companies = len({job.company for job in result.jobs})
            st.session_state.aja_search_report = {
                "mode": "local-playwright",
                "boards_queried": companies,
                "boards_succeeded": companies,
                "total_jobs_seen": len(result.jobs),
                "matched_jobs": len(result.jobs),
                "warnings": result.warnings,
                "exported_at": result.exported_at,
            }
            st.session_state.aja_notice = f"已导入 {len(result.jobs)} 个岗位并完成匹配；没有自动访问或提交任何招聘页面。"
            st.rerun()
        except (LocalJobImportError, ValueError) as exc:
            _clear_job_search_state()
            st.error(str(exc))


def _search_step(profile: CandidateProfile, service: ApplicationService) -> None:
    st.markdown("### 3 · 找到匹配岗位")
    source_label = st.radio(
        "岗位来源",
        (
            "国内企业 · 校园招聘",
            "外企 / 国际企业 · 校招与初级岗位",
            "指定企业官网 / 其他 ATS",
            "导入本地 Playwright 岗位",
            "本地 Demo（4 个样例）",
        ),
        horizontal=True,
        key="aja_job_source",
        on_change=_clear_job_search_state,
        help="最低匹配度只负责筛选已找到的岗位，不会增加公司来源。",
    )
    if source_label == "导入本地 Playwright 岗位":
        _local_jobs_step(profile, service)
        return
    use_demo_jobs = source_label.startswith("本地 Demo")
    use_domestic_sources = source_label.startswith("国内企业")
    use_international_sources = source_label.startswith("外企 / 国际企业")
    use_campus_sources = use_domestic_sources
    use_custom_sources = source_label.startswith("指定企业官网")
    use_official_url_sources = use_campus_sources or use_custom_sources
    search_scope = (
        "domestic"
        if use_domestic_sources
        else ("international" if use_international_sources else None)
    )
    official_urls: list[str] = []
    official_company_by_url: dict[str, str] = {}

    if use_domestic_sources:
        all_regions = tuple(
            dict.fromkeys(region for site in CAMPUS_CAREER_SITES for region in site.regions)
        )
        selected_regions = set(
            st.multiselect(
                "校招地区",
                all_regions,
                default=["中国大陆"] if "中国大陆" in all_regions else list(all_regions[:1]),
                help="这里只改变可选公司目录；岗位地点仍会在结果中单独显示。",
            )
        )
        domestic_labels = {
            site.label for site in campus_sites_for_scope("domestic")
        }
        available_sites = tuple(
            site
            for site in campus_sites_for_regions(selected_regions)
            if site.label in domestic_labels
        )
        searchable_sites = [
            site for site in available_sites if site.discovery_mode != "link_only"
        ]
        automatic_sites = [
            site for site in available_sites if site.discovery_mode == "automatic"
        ]
        available_labels = [site.label for site in searchable_sites]
        default_labels = [
            label for label in DEFAULT_CAMPUS_LABELS if label in available_labels
        ]
        automatic_labels = [site.label for site in automatic_sites]
        range_mode = st.radio(
            "搜索范围",
            (
                f"广泛搜索（推荐 · {len(default_labels)} 家）",
                f"快速搜索（已验证 · {len(automatic_labels)} 家）",
                "自定义公司",
            ),
            horizontal=True,
            help="广泛搜索覆盖更多企业；快速搜索只读取当前已取得岗位结果的渠道。",
        )
        if range_mode.startswith("广泛搜索"):
            presets = default_labels
        elif range_mode.startswith("快速搜索"):
            presets = automatic_labels
        else:
            presets = st.multiselect(
                "选择官方校招渠道",
                available_labels,
                default=automatic_labels,
                max_selections=MAX_CAMPUS_SOURCES_PER_SEARCH,
                placeholder="选择本轮要搜索的公司",
                help=(
                    "为尊重网站限制并控制等待时间，每轮最多 "
                    f"{MAX_CAMPUS_SOURCES_PER_SEARCH} 家。"
                ),
            )
        selected_sites = resolve_campus_selection(presets)
        for site in selected_sites:
            official_urls.append(site.url)
            official_company_by_url[site.url] = site.company
        with st.expander(f"查看 {len(available_sites)} 个国内企业官方校招入口"):
            for site in available_sites:
                mode_text = {
                    "automatic": "✓ 当前可自动读取岗位",
                    "best_effort": "动态页面，尽力读取",
                    "link_only": "官网可直接打开，专用适配器待接入",
                }[site.discovery_mode]
                st.markdown(
                    f"- [{site.label}]({site.url}) · {site.system} · "
                    f"{mode_text}"
                )
        st.caption(
            f"国内企业目录共 {len(available_sites)} 家；本轮可尝试读取 "
            f"{len(searchable_sites)} 家，其中 {len(automatic_sites)} 家已取得真实岗位结果。"
            "动态站若未返回岗位，仍可从上方打开官方入口。"
        )
        st.caption(f"本轮已选择 {len(selected_sites)} 家官方渠道。")
        c1, c2 = st.columns([1, 1.7])
        result_limit = int(c1.selectbox("每家公司最多读取", (25, 50, 100), index=1))
        ai_review = c2.checkbox(
            "使用 AI 复核排名前 30 个岗位（可选）",
            value=False,
            disabled=not SETTINGS.openai_enabled,
        )
        early_career_only = st.checkbox(
            "隐藏明显属于社招、资深或要求多年经验的岗位",
            value=True,
            help="校招页面上没有写“应届”的普通岗位仍会保留；明确资深岗位会隐藏。",
        )
        broad_search = True
        nationwide = True
        search_locations = ()
        st.caption(
            "本区只显示国内企业，按企业属性分组而不是按岗位地点分组。"
            "动态官网若无法自动提取岗位，会保留官方入口供你直接打开；"
            "不会绕过登录、验证码、robots 或网站限制。"
        )
    elif use_demo_jobs:
        broad_search = False
        nationwide = False
        search_locations = profile.target_locations
        result_limit = 20
        ai_review = False
        early_career_only = False
        st.caption("本地 Demo 永远只有 4 个固定样例，用于验证界面和安全投递流程。")
    elif use_custom_sources:
        presets = st.multiselect(
            "选择已核验的企业招聘入口（可多选）",
            list(CAREER_SITE_BY_LABEL),
            placeholder="例如：西门子、字节跳动、小米",
            help="自建系统没有统一 API；程序会在公开、同域、有限页数内读取。",
        )
        custom_urls = st.text_area(
            "或粘贴企业官方招聘页 / Lever / Ashby / SmartRecruiters 地址",
            placeholder=(
                "每行一个，例如：\n"
                "https://jobs.lever.co/company\n"
                "https://jobs.ashbyhq.com/company\n"
                "https://careers.company.com/jobs"
            ),
            height=110,
        )
        for preset_label in presets:
            site = CAREER_SITE_BY_LABEL[preset_label]
            official_urls.append(site.url)
            official_company_by_url[site.url] = site.company
        official_urls.extend(
            line.strip() for line in custom_urls.splitlines() if line.strip()
        )
        c1, c2 = st.columns([1, 1.6])
        result_limit = int(c1.selectbox("每个来源最多读取", (50, 100, 200), index=1))
        ai_review = c2.checkbox(
            "使用 AI 复核排名前 30 个岗位（可选）",
            value=False,
            disabled=not SETTINGS.openai_enabled,
        )
        broad_search = True
        nationwide = True
        search_locations = ()
        early_career_only = st.checkbox(
            "只保留校招 / 应届 / 实习岗位",
            value=profile.career_stage in {"fresh_graduate", "internship"},
        )
        st.caption(
            "自动识别 Lever、Ashby、SmartRecruiters；其他自建官网使用 "
            f"JobPosting / 岗位卡片通用识别。最多 {MAX_CAMPUS_SOURCES_PER_SEARCH} 个来源，"
            "robots、登录、验证码或限流会暂停。"
        )
        if SETTINGS.demo_mode:
            st.info("当前是安全 Demo Mode：可读取真实公开岗位；官网预填需切换到 Real Mode。")
    else:
        international_campus_sites = campus_sites_for_scope("international")
        with st.expander(
            f"查看 {len(international_campus_sites)} 个外企中国校招官网"
        ):
            for site in international_campus_sites:
                st.markdown(
                    f"- [{site.label}]({site.url}) · {site.system} · 官网直接打开"
                )
        c1, c2, c3 = st.columns([1.2, 1.35, 0.85])
        recall_mode = c1.selectbox(
            "搜索方式",
            ("宽泛推荐", "严格匹配岗位名称"),
            help="宽泛推荐先读取更多岗位，再由匹配算法排序。",
        )
        location_scope = c2.selectbox(
            "搜索地区",
            ("中国大陆 + 香港 + 新加坡", "只看目标城市", "不限城市（当前公司范围）"),
        )
        result_limit = int(c3.selectbox("最多读取", (100, 200, 500), index=1))
        broad_search = recall_mode == "宽泛推荐"
        if location_scope == "中国大陆 + 香港 + 新加坡":
            search_locations = ("中国", "香港", "新加坡")
            nationwide = False
        elif location_scope == "只看目标城市":
            search_locations = profile.target_locations
            nationwide = not bool(search_locations)
        else:
            search_locations = ()
            nationwide = True
        ai_review = st.checkbox(
            "使用 AI 复核排名前 30 个岗位（可选）",
            value=False,
            disabled=not SETTINGS.openai_enabled,
            help="未配置 OPENAI_API_KEY 时使用固定权重规则；不会为数百个岗位逐个调用 API。",
        )
        early_career_only = st.checkbox(
            "只保留校招 / 应届 / 实习岗位",
            value=profile.career_stage in {"fresh_graduate", "internship"},
            help="应届生模式默认开启；关闭后也可查看该公司范围内的全部岗位。",
        )
        st.caption(
            f"本区将检查 {len(DEFAULT_GREENHOUSE_TOKENS)} 个经核验的国际职位来源，"
            "再保留校招、毕业生和实习岗位；其中 5 个企业主体口径不清的来源不会混入外企结果。"
            "公司失败会被隔离，缓存有效期 6 小时。"
        )
        if SETTINGS.demo_mode:
            st.info("当前是安全 Demo Mode：可以搜索真实公开岗位，但投递动作仍只在本地模拟。")

    label = (
        "加载 4 个本地样例"
        if use_demo_jobs
        else (
            f"读取 {len(set(official_urls))} 个指定官网"
            if use_custom_sources
            else (
                f"搜索 {len(set(official_urls))} 个国内企业校招渠道"
                if use_domestic_sources
                else f"搜索 {len(DEFAULT_GREENHOUSE_TOKENS)} 个国际职位来源"
            )
        )
    )
    if st.button(
        label,
        type="primary",
        use_container_width=True,
        disabled=use_official_url_sources and not official_urls,
    ):
        try:
            with st.spinner("正在发现岗位并按固定权重计算匹配度…"):
                if use_official_url_sources:
                    official_result = search_official_career_urls(
                        official_urls,
                        company_by_url=official_company_by_url,
                        max_jobs_per_source=result_limit,
                    )
                    jobs = list(official_result.jobs)
                else:
                    result = search_jobs_detailed(
                        profile,
                        demo_mode=use_demo_jobs,
                        target_roles=profile.target_roles,
                        target_locations=search_locations,
                        broad_search=broad_search,
                        nationwide=nationwide,
                        limit=result_limit,
                    )
                    jobs = list(result.jobs)
                discovered_job_count = len(jobs)
                jobs = filter_jobs_by_scope(jobs, search_scope)
                scope_filtered_out = discovered_job_count - len(jobs)
                unfiltered_job_count = len(jobs)
                if early_career_only:
                    jobs = filter_early_career_jobs(
                        jobs,
                        verified_campus_channel=use_campus_sources,
                    )
                filtered_out = unfiltered_job_count - len(jobs)
                assessments: dict[str, MatchAssessment] = {}
                for job in jobs:
                    assessment = match_job(
                        profile,
                        job,
                        use_openai=False,
                        model=SETTINGS.openai_model,
                    )
                    assessments[_job_key(job)] = assessment
                    service.register_job(job, assessment)
                if ai_review and jobs:
                    ranked = sorted(
                        jobs,
                        key=lambda item: assessments[_job_key(item)].match_score,
                        reverse=True,
                    )[:30]
                    for job in ranked:
                        assessment = match_job(
                            profile,
                            job,
                            use_openai=True,
                            model=SETTINGS.openai_model,
                        )
                        assessments[_job_key(job)] = assessment
                        service.register_job(job, assessment)
            st.session_state.aja_jobs = jobs
            st.session_state.aja_assessments = assessments
            st.session_state.aja_selected = set()
            if use_official_url_sources:
                st.session_state.aja_search_report = {
                    "mode": "domestic-career" if use_domestic_sources else "official-career",
                    "company_scope": search_scope,
                    "boards_queried": len(official_result.sources_queried),
                    "boards_succeeded": len(official_result.sources_succeeded),
                    "failures": list(official_result.failures),
                    "warnings": list(official_result.warnings),
                    "cache_hits": 0,
                    "total_jobs_seen": official_result.total_jobs_seen,
                    "matched_jobs": len(jobs),
                    "filtered_out": filtered_out,
                    "scope_filtered_out": scope_filtered_out,
                    "pages_fetched": official_result.pages_fetched,
                }
            else:
                st.session_state.aja_search_report = {
                    "mode": "international-career" if use_international_sources else result.mode,
                    "company_scope": search_scope,
                    "boards_queried": len(result.boards_queried),
                    "boards_succeeded": len(result.boards_succeeded),
                    "failures": list(result.failures),
                    "cache_hits": len(result.cache_hits),
                    "total_jobs_seen": result.total_jobs_seen,
                    "matched_jobs": len(jobs),
                    "filtered_out": filtered_out,
                    "scope_filtered_out": scope_filtered_out,
                }
            st.rerun()
        except (JobSearchError, OfficialCareerSearchError) as exc:
            st.error(str(exc))


def _job_filters(
    jobs: list[JobPosting],
) -> tuple[float, set[str], set[str], set[str], set[str]]:
    c1, c2, c3, c4, c5 = st.columns([1, 1.15, 1.2, 1.2, 1.2])
    minimum = c1.slider(
        "最低匹配度（仅显示）",
        0,
        100,
        70,
        5,
        help="降低它只会显示更多已读取岗位，不会搜索更多公司。",
    )
    company_scopes = c2.multiselect(
        "企业类型",
        sorted({scope_label(company_scope_for_job(job)) for job in jobs}),
    )
    companies = c3.multiselect("公司", sorted({job.company for job in jobs}))
    locations = c4.multiselect("地点", sorted({job.location or "未注明" for job in jobs}))
    types = c5.multiselect("岗位类型", sorted({job.job_type or "未注明" for job in jobs}))
    return (
        float(minimum),
        set(company_scopes),
        set(companies),
        set(locations),
        set(types),
    )


def _similar_jobs(job: JobPosting, visible: list[tuple[JobPosting, MatchAssessment]]) -> list[str]:
    candidates: list[tuple[float, str]] = []
    current_words = set(job.title.casefold().split())
    for other, assessment in visible:
        if _job_key(other) == _job_key(job):
            continue
        words = set(other.title.casefold().split())
        overlap = len(current_words & words)
        same_type = int(bool(job.job_type and job.job_type == other.job_type))
        candidates.append((overlap * 10 + same_type * 3 + assessment.match_score / 100, f"{other.company} · {other.title}"))
    return [label for _, label in sorted(candidates, reverse=True)[:2]]


def _jobs_list(service: ApplicationService) -> None:
    raw_jobs = st.session_state.aja_jobs
    jobs = [item if isinstance(item, JobPosting) else JobPosting.model_validate(item) for item in raw_jobs]
    report = dict(st.session_state.aja_search_report or {})
    if report:
        if report.get("mode") == "local-playwright":
            st.caption("来源：本机 Playwright 导出 · 导入匹配不请求招聘网站；请在投递前核对岗位是否仍然开放。")
        coverage = st.columns(5)
        coverage[0].metric("读取公司", report.get("boards_succeeded", 0))
        coverage[1].metric("官网岗位池", report.get("total_jobs_seen", len(jobs)))
        coverage[2].metric("进入匹配", report.get("matched_jobs", len(jobs)))
        coverage[3].metric("隐藏社招", report.get("filtered_out", 0))
        coverage[4].metric("缓存来源", report.get("cache_hits", 0))
        failures = list(report.get("failures") or [])
        if failures:
            with st.expander(f"{len(failures)} 个公司暂时读取失败（不影响其他结果）"):
                for failure in failures:
                    st.write(f"- {failure}")
        warnings = list(report.get("warnings") or [])
        if warnings:
            with st.expander(f"{len(warnings)} 条官网读取提示"):
                for warning in warnings:
                    st.write(f"- {warning}")
        scope_filtered_out = int(report.get("scope_filtered_out") or 0)
        if scope_filtered_out:
            st.caption(f"企业类型校验已排除 {scope_filtered_out} 个不属于当前分区的岗位。")

    if not jobs:
        if report:
            st.info(
                "本轮已检查所选来源，但暂未提取到同时符合企业类型和应届条件的岗位。"
                "你仍可展开上方官方渠道并直接打开官网；动态页面后续会继续增加专用适配器。"
            )
        return

    assessments: dict[str, MatchAssessment] = {}
    for key, value in st.session_state.aja_assessments.items():
        assessments[key] = value if isinstance(value, MatchAssessment) else MatchAssessment.model_validate(value)

    scope_counts = Counter(scope_label(company_scope_for_job(job)) for job in jobs)
    st.caption(
        f"国内企业 {scope_counts['国内企业']} 个 · "
        f"外企 / 国际企业 {scope_counts['外企 / 国际企业']} 个 · "
        f"未分类 {scope_counts['未分类']} 个"
    )
    minimum, company_scopes, companies, locations, types = _job_filters(jobs)
    visible = [
        (job, assessments[_job_key(job)])
        for job in jobs
        if assessments[_job_key(job)].match_score >= minimum
        and (
            not company_scopes
            or scope_label(company_scope_for_job(job)) in company_scopes
        )
        and (not companies or job.company in companies)
        and (not locations or (job.location or "未注明") in locations)
        and (not types or (job.job_type or "未注明") in types)
    ]
    visible.sort(key=lambda item: item[1].match_score, reverse=True)
    st.caption(
        f"符合当前显示条件：{len(visible)} 个岗位 · "
        "分数由固定权重规则计算，AI 启用时最多只做 ±5 分复核"
    )
    if not visible:
        st.info("当前显示条件下没有岗位。可降低最低匹配度，或在上方增加更多官方招聘来源。")
        return

    page_size = int(st.selectbox("每页显示", (20, 50, 100), index=0, key="aja_page_size"))
    page_count = max(1, (len(visible) + page_size - 1) // page_size)
    if int(st.session_state.get("aja_page_number", 1)) > page_count:
        st.session_state.aja_page_number = 1
    page_number = int(
        st.number_input(
            "页码",
            min_value=1,
            max_value=page_count,
            value=1,
            step=1,
            key="aja_page_number",
        )
    )
    start = (page_number - 1) * page_size
    page_jobs = visible[start : start + page_size]
    st.caption(f"第 {page_number} / {page_count} 页，当前显示 {len(page_jobs)} 个")

    selected = set(st.session_state.aja_selected)
    for job, assessment in page_jobs:
        key = _job_key(job)
        with st.container(border=True):
            pick, body, score = st.columns([0.09, 0.72, 0.19], vertical_alignment="top")
            checked = pick.checkbox(
                "选择",
                value=key in selected,
                key=f"aja-pick-{key}",
                label_visibility="collapsed",
            )
            if checked:
                selected.add(key)
            else:
                selected.discard(key)
            body.markdown(f"#### {markdown_literal(job.company)} · {markdown_literal(job.title)}")
            body.caption(
                markdown_literal(" · ".join(
                    filter(
                        None,
                        [
                            scope_label(company_scope_for_job(job)),
                            "官方校招" if report.get("mode") == "domestic-career" else None,
                            job.location,
                            job.job_type,
                            job.department,
                            job.source,
                        ],
                    )
                ))
            )
            if assessment.advantages:
                body.write("**优势**　" + markdown_literal("；".join(assessment.advantages[:2])))
            if assessment.missing_skills:
                body.caption("缺失技能：" + "、".join(assessment.missing_skills[:6]))
            score.markdown(f'<div class="aja-score">{assessment.match_score:.0f}%</div>', unsafe_allow_html=True)
            score.caption(assessment.match_level or "")
            with st.expander("查看 JD、评分依据与同类岗位"):
                st.text(job.description or "暂无 JD 正文")
                if job.requirements:
                    st.markdown("**要求**")
                    st.text("\n".join(job.requirements) if isinstance(job.requirements, list) else job.requirements)
                st.markdown("**评分依据**")
                st.text(assessment.reason)
                similar = _similar_jobs(job, visible)
                st.caption("同类型岗位：" + markdown_literal("；".join(similar) if similar else "暂无"))
                if job.job_url.startswith("https://"):
                    st.link_button("打开官网岗位页", job.job_url)

    st.session_state.aja_selected = selected
    st.caption(f"已选择 {len(selected)} / 20 个岗位")
    if len(selected) > 20:
        st.error("第一版一次最多准备 20 个岗位，请取消部分选择。")
        return
    if SETTINGS.cloud_deployment:
        action_label = "加入云端投递清单"
        st.caption(
            "云端版只生成岗位清单和官网入口，不会启动浏览器、发送简历或提交申请。"
        )
    elif SETTINGS.demo_mode:
        action_label = "一键准备投递"
        st.caption("Demo Mode 只在本机模拟表单，不会向外部网站发送数据。")
    else:
        action_label = "一键投递：自动打开官网并预填"
        st.caption(
            "正式模式会按顺序逐个打开官网：已知资料自动填写，未知问题保持空白，"
            "最终提交由你本人点击。"
        )
    if st.button(
        action_label,
        type="primary",
        use_container_width=True,
        disabled=not selected,
    ):
        _prepare_queue(service, jobs, assessments, selected)
        st.rerun()


def _prepare_queue(
    service: ApplicationService,
    jobs: list[JobPosting],
    assessments: dict[str, MatchAssessment],
    selected: set[str],
) -> None:
    profile = _profile_from_state()
    if profile is None:
        st.error("请先解析简历。")
        return
    values = _profile_values(profile)
    created = 0
    existing = 0
    active = 0
    plan_failures = 0
    launch_requests: list[AutofillLaunchRequest] = []
    demo_result: BrowserRunResult | None = None
    if SETTINGS.demo_mode and not SETTINGS.cloud_deployment:
        demo_result = _run_local_demo_browser(values, confirmed=False)
    for job in jobs:
        key = _job_key(job)
        if key not in selected:
            continue
        try:
            queue = service.enqueue_job(
                job,
                assessment=assessments[key],
                resume_version=Path(values["resume"]).name,
            )
            queue_id = int(queue["id"])
            if queue.get("already_exists"):
                existing += 1
            else:
                created += 1
            if SETTINGS.cloud_deployment:
                continue
            if SETTINGS.demo_mode:
                if queue.get("already_exists"):
                    continue
                service.transition(queue_id, "opening")
                service.transition(queue_id, "filling")
                assert demo_result is not None
                if demo_result.status == "error" or demo_result.review is None:
                    metadata = demo_result.screenshot_metadata or {}
                    service.fail(
                        queue_id,
                        "Playwright Demo 表单填写失败",
                        screenshot=str(metadata.get("screenshot_path") or ""),
                    )
                    continue
                browser_review = demo_result.review
                public_review = demo_result.public_summary()
                confirmations = [item.public_summary() for item in browser_review.pending_confirmations]
                blocked = bool(browser_review.blockers)
                service.record_review(
                    queue_id,
                    public_review,
                    needs_user_confirmation=confirmations,
                    required_field_missing=blocked or bool(confirmations),
                )
                service.transition(
                    queue_id,
                    "waiting_user" if blocked or confirmations else "ready_to_submit",
                )
                continue

            status = str(queue.get("status") or "pending")
            if status == "failed":
                service.retry_failed(queue_id)
                queue["status"] = "pending"
                status = "pending"
            if status in {"submitted", "skipped"}:
                continue
            if status in {"opening", "filling"} or is_autofill_scheduled(queue_id):
                active += 1
                continue
            try:
                launch_requests.append(
                    _build_official_autofill_request(service, queue, profile)
                )
            except (AutofillError, KeyError, TypeError, ValueError) as exc:
                service.fail(queue_id, str(exc) or "无法创建官网填写计划")
                plan_failures += 1
        except DuplicateApplicationError:
            existing += 1

    if SETTINGS.cloud_deployment:
        st.session_state.aja_notice = (
            f"已加入 {created} 个岗位到云端清单；{existing} 个已有岗位未重复创建。"
            "云端不会启动浏览器或发送个人资料。"
        )
        return

    if SETTINGS.demo_mode:
        st.session_state.aja_notice = (
            f"已准备 {created} 个岗位；{existing} 个已有岗位未重复创建。"
        )
        return

    handle = schedule_autofill_batch(service, launch_requests)
    scheduled = len(handle.scheduled_queue_ids)
    message = (
        f"已新增 {created} 个岗位，并启动 {scheduled} 个官网投递任务。"
        "多选岗位会逐个打开：完成当前官网并关闭窗口后，下一家会自动打开。"
    )
    if existing:
        message += f" {existing} 个已有岗位已复用原队列。"
    if active:
        message += f" {active} 个岗位已经在打开或填写中。"
    if plan_failures:
        message += f" {plan_failures} 个岗位无法创建安全填写计划，已标记为失败供你重试。"
    st.session_state.aja_notice = message


def _render_real_submission_confirmation(
    service: ApplicationService, item: Mapping[str, Any]
) -> None:
    queue_id = int(item["id"])
    submitted_ack = st.checkbox(
        "我已在官网补全空白、核对并亲自提交",
        key=f"aja-submitted-ack-{queue_id}",
    )
    if st.button(
        "保存为已投递",
        key=f"aja-record-submitted-{queue_id}",
        disabled=not submitted_ack,
        use_container_width=True,
    ):
        service.record_review(
            queue_id,
            {
                "mode": "user_verified_official_submission",
                "user_attested_final_review": True,
            },
            needs_user_confirmation=[],
            required_field_missing=False,
        )
        if str(item["status"]) != "ready_to_submit":
            service.transition(queue_id, "ready_to_submit")
        service.confirm_submission(queue_id, user_confirmed=True)
        st.session_state.aja_notice = "真实投递记录已保存。"
        st.rerun()


@st.fragment(run_every=2.0)
def _queue_section(service: ApplicationService) -> None:
    queue = service.list_queue(limit=100)
    if not queue:
        return
    st.divider()
    st.markdown("### 4 · 投递队列")
    if st.session_state.aja_notice:
        st.success(st.session_state.aja_notice)
        st.session_state.aja_notice = ""

    for item in queue:
        status = str(item["status"])
        with st.container(border=True):
            left, middle, right = st.columns([1.6, 1, 1])
            left.write(f"**{markdown_literal(item['company'])} · {markdown_literal(item['job_title'])}**")
            left.caption(f"匹配度 {item.get('match_score') or '—'} · {STATUS_LABELS.get(status, status)}")
            if SETTINGS.cloud_deployment:
                middle.info("云端投递清单")
                job_url = str(item.get("job_url") or "")
                if job_url.startswith("https://"):
                    right.link_button("打开岗位官网", job_url)
                st.caption("本记录保存在临时云实例中；实例重启后可能清除。")
                continue
            if status == "ready_to_submit":
                if bool(item.get("is_demo")):
                    middle.success("✓ 表单检查通过")
                    if right.button(
                        "确认模拟提交",
                        type="primary",
                        key=f"aja-submit-{item['id']}",
                    ):
                        profile = _profile_from_state()
                        if profile is None:
                            st.error("候选人档案已失效，请重新加载简历。")
                        else:
                            result = _run_local_demo_browser(
                                _profile_values(profile), confirmed=True
                            )
                            if result.submitted_demo:
                                service.confirm_submission(int(item["id"]), user_confirmed=True)
                                st.session_state.aja_notice = "Demo 投递已由 Playwright 在本地模拟完成并保存记录；没有向外部网站发送数据。"
                                st.rerun()
                            else:
                                st.error("本地 Demo 未通过提交门禁，未保存为已投递。")
                else:
                    review = dict(item.get("review") or {})
                    filled = list(review.get("filled_fields") or [])
                    middle.success("官网已预填完成")
                    right.caption(f"已自动填写 {len(filled)} 项" if filled else "等待你核对")
                    st.info(
                        "官网窗口已经打开并完成安全预填。请核对页面，补充仍为空的项目，"
                        "然后由你本人点击官网的最终提交按钮。"
                    )
                    _render_real_submission_confirmation(service, item)
            elif status == "waiting_user":
                review = dict(item.get("review") or {})
                visible_browser = review.get("mode") == "visible_official_browser"
                filled = list(review.get("filled_fields") or [])
                middle.warning("官网已打开，空白待补充" if visible_browser else "等待你在官网继续")
                if visible_browser:
                    right.caption(f"已自动填写 {len(filled)} 项" if filled else "正在等待官网操作")
                needs = item.get("needs_user_confirmation") or []
                with st.expander(
                    "这些官网项目已保持空白，请你填写" if visible_browser else "查看待确认项目",
                    expanded=True,
                ):
                    if needs:
                        for pending in needs:
                            label = pending.get("field_name", "未知字段") if isinstance(pending, dict) else str(pending)
                            reason = pending.get("reason", "需要本人确认") if isinstance(pending, dict) else "需要本人确认"
                            st.write(f"- **{label}**：{reason}")
                    else:
                        st.write("请回到已打开的官网窗口继续；程序不会猜测或代填未知答案。")
                    st.caption("你的填写内容不会被读取或写入日志，最终提交仍由你本人完成。")
                if not bool(item.get("is_demo")):
                    profile = _profile_from_state()
                    if st.button(
                        "重新打开并预填",
                        key=f"aja-reopen-{item['id']}",
                        disabled=profile is None or is_autofill_scheduled(int(item["id"])),
                    ):
                        assert profile is not None
                        try:
                            started = _start_visible_official_application(service, item, profile)
                            st.session_state.aja_notice = (
                                "已重新加入自动打开队列；程序会填写已知资料并保留未知空白。"
                                if started
                                else "该岗位已经在自动打开队列中。"
                            )
                            st.rerun()
                        except AutofillError as exc:
                            st.error(str(exc))
                        except Exception:
                            st.error("官网预填未能启动；没有提交任何申请。")
                    _render_real_submission_confirmation(service, item)
            elif status == "failed":
                middle.error(item.get("last_error") or "自动填写失败")
                if right.button("重试", key=f"aja-retry-{item['id']}"):
                    service.retry_failed(int(item["id"]))
                    st.rerun()
            elif status == "submitted":
                middle.success("已保存投递记录")
                right.caption("Demo 模拟" if bool(item.get("is_demo")) else "真实投递")
            else:
                middle.info(STATUS_LABELS.get(status, status))
                if not SETTINGS.demo_mode and status == "pending":
                    profile = _profile_from_state()
                    scheduled = is_autofill_scheduled(int(item["id"]))
                    if scheduled:
                        right.caption("等待前一个官网完成")
                    if right.button(
                        "立即加入自动投递",
                        type="primary",
                        key=f"aja-open-fill-{item['id']}",
                        disabled=profile is None or scheduled,
                    ):
                        assert profile is not None
                        try:
                            started = _start_visible_official_application(service, item, profile)
                            st.session_state.aja_notice = (
                                "已加入自动打开队列：已知资料会自动填写，未知问题保持空白。"
                                if started
                                else "该岗位已经在自动打开队列中。"
                            )
                            st.rerun()
                        except AutofillError as exc:
                            st.error(str(exc))
                        except Exception:
                            st.error("官网预填未能启动；没有提交任何申请。")
                    left.link_button("只打开官网", str(item["job_url"]))


def render_agent_home() -> None:
    """Render the simple primary product flow inside the first Streamlit tab."""

    _init_state()
    _restore_saved_candidate_profile()
    _inject_style()
    service = _application_service()
    if not st.session_state.aja_legacy_queue_migrated:
        reset_count = service.reset_legacy_autofill_placeholders()
        st.session_state.aja_legacy_queue_migrated = True
        if reset_count:
            st.session_state.aja_notice = (
                f"已清除 {reset_count} 个旧版笼统“人工处理”占位状态；"
                "现在可直接一键打开官网并按真实空白项处理。"
            )
    mode_label = (
        "CLOUD SAFE MODE"
        if SETTINGS.cloud_deployment
        else "DEMO MODE" if SETTINGS.demo_mode else "REAL MODE"
    )
    st.markdown(
        f"""
        <section class="aja-hero">
          <div class="aja-eyebrow">{mode_label}</div>
          <h1>{APP_NAME}</h1>
          <p>上传一次简历，自动解析、发现岗位、计算匹配度并准备官网表单。
          <span class="aja-safe">未知信息不猜，验证码不绕过，最终提交由你确认。</span></p>
          <div class="aja-flow">上传简历　→　AI 解析　→　匹配岗位　→　准备投递　→　确认提交</div>
        </section>
        """,
        unsafe_allow_html=True,
    )
    _dashboard(service)
    st.divider()
    _resume_step()
    profile = _profile_from_state()
    if profile is None:
        if SETTINGS.demo_mode:
            st.caption("没有 API Key 也可以点击“使用示例简历”体验完整流程。")
        else:
            st.caption("正式模式已启用：请上传并解析你的真实 PDF / DOCX 简历后开始。")
        _queue_section(service)
        return
    profile = _profile_step(profile)
    _search_step(profile, service)
    _jobs_list(service)
    _queue_section(service)


__all__ = ["render_agent_home"]
