from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import BinaryIO

MAX_RESUME_BYTES = 10 * 1024 * 1024


class ResumeError(ValueError):
    pass


def extract_resume_text(filename: str, content: bytes) -> str:
    if not content:
        raise ResumeError("简历文件是空的")
    if len(content) > MAX_RESUME_BYTES:
        raise ResumeError("简历文件不能超过 10 MB")

    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        text = _decode_text(content)
    elif suffix == ".pdf":
        text = _read_pdf(BytesIO(content))
    elif suffix == ".docx":
        text = _read_docx(BytesIO(content))
    else:
        raise ResumeError("仅支持 PDF、DOCX、TXT 或 Markdown 简历")

    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) < 20:
        raise ResumeError("没有从简历中读到足够文本；扫描版 PDF 请先做 OCR")
    return cleaned


def read_resume_file(path: str | Path) -> str:
    resume_path = Path(path)
    return extract_resume_text(resume_path.name, resume_path.read_bytes())


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResumeError("文本简历编码无法识别，请保存为 UTF-8")


def _read_pdf(stream: BinaryIO) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ResumeError("读取 PDF 需要安装 pypdf") from exc
    try:
        return "\n".join(page.extract_text() or "" for page in PdfReader(stream).pages)
    except Exception as exc:
        raise ResumeError(f"PDF 读取失败：{exc}") from exc


def _read_docx(stream: BinaryIO) -> str:
    try:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ResumeError("读取 DOCX 需要安装 python-docx") from exc
    try:
        document = Document(stream)
        lines: list[str] = []

        def add(value: str) -> None:
            cleaned = " ".join(str(value or "").split())
            if cleaned and (not lines or lines[-1] != cleaned):
                lines.append(cleaned)

        # 很多简历使用无边框表格排版，需要按正文顺序读取段落和表格。
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                add(Paragraph(child, document).text)
            elif isinstance(child, CT_Tbl):
                table = Table(child, document)
                for row in table.rows:
                    for cell in row.cells:
                        add(cell.text)

        # 联系方式和页码经常放在页眉或页脚中。
        for section in document.sections:
            for paragraph in section.header.paragraphs:
                add(paragraph.text)
            for table in section.header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        add(cell.text)
            for paragraph in section.footer.paragraphs:
                add(paragraph.text)

        # python-docx 不把文本框暴露为普通段落，单独补取其中的文字。
        for text_node in document.part.element.xpath(".//w:txbxContent//w:t"):
            add(text_node.text or "")
        return "\n".join(lines)
    except Exception as exc:
        raise ResumeError(f"DOCX 读取失败：{exc}") from exc
