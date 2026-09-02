from __future__ import annotations

from io import BytesIO


class ResumeExportError(ValueError):
    """A tailored resume draft could not be exported safely."""


SECTION_HEADINGS = {
    "个人信息", "基本信息", "联系方式", "求职目标", "求职意向", "个人简介",
    "自我评价", "工作经历", "实习经历", "项目经历", "项目经验", "教育经历",
    "教育背景", "技能", "专业技能", "证书", "获奖经历", "校园经历", "语言能力",
    "contact", "contact information", "objective", "summary", "profile",
    "experience", "work experience", "projects", "project experience", "education",
    "skills", "technical skills", "certifications", "awards", "languages",
}


def resume_text_to_docx(text: str) -> bytes:
    """Render an approved plain-text resume draft as a clean DOCX.

    The function is deliberately presentation-only: it does not rewrite, add,
    translate, or remove any resume fact.
    """

    cleaned_lines = [line.strip() for line in str(text or "").splitlines() if line.strip()]
    if not cleaned_lines:
        raise ResumeExportError("定制简历正文不能为空")
    if len("\n".join(cleaned_lines)) > 200_000:
        raise ResumeExportError("定制简历正文过长")

    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise ResumeExportError("导出 DOCX 需要安装 python-docx") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.start_type = WD_SECTION.NEW_PAGE

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    first_content = True
    for line in cleaned_lines:
        normalized_heading = line.lstrip("#").strip().rstrip(":：").casefold()
        if normalized_heading in SECTION_HEADINGS:
            paragraph = document.add_heading(line.lstrip("#").strip().rstrip(":："), level=1)
            paragraph.paragraph_format.space_before = Pt(7)
            paragraph.paragraph_format.space_after = Pt(2)
        elif line.startswith(("- ", "• ", "· ", "* ")):
            paragraph = document.add_paragraph(line[2:].strip(), style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(0)
        elif first_content:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line)
            run.bold = True
            run.font.size = Pt(16)
        else:
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(1.5)
        first_content = False

    output = BytesIO()
    try:
        document.save(output)
    except Exception as exc:  # pragma: no cover - python-docx normally wraps these
        raise ResumeExportError(f"定制简历导出失败：{exc}") from exc
    return output.getvalue()


__all__ = ["ResumeExportError", "resume_text_to_docx"]
