from io import BytesIO

import pytest
from docx import Document

from job_assistant.resume import ResumeError, extract_resume_text


def test_reads_utf8_text_resume():
    text = extract_resume_text("resume.txt", "姓名：测试\n技能：Python、SQL\n项目：完成岗位匹配系统".encode())
    assert "Python" in text
    assert "岗位匹配" in text


def test_rejects_unsupported_resume_type():
    with pytest.raises(ResumeError, match="仅支持"):
        extract_resume_text("resume.exe", b"not a resume but definitely long enough")


def test_docx_reader_includes_tables_and_header_content():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "candidate@example.com"
    document.add_paragraph("智能制造工程")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python 数据分析"
    table.cell(0, 1).text = "SolidWorks"
    stream = BytesIO()
    document.save(stream)

    text = extract_resume_text("resume.docx", stream.getvalue())

    assert "智能制造工程" in text
    assert "Python 数据分析" in text
    assert "SolidWorks" in text
    assert "candidate@example.com" in text
