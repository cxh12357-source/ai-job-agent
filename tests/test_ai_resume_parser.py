from io import BytesIO

import pymupdf
import pytest
from docx import Document

from ai_job_agent.services.resume_parser import ResumeParseError, parse_resume


def test_ai_resume_parser_reads_pdf():
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Candidate Python Data Analysis Experience")
    payload = document.tobytes()
    document.close()

    assert "Python Data Analysis" in parse_resume("resume.pdf", payload)


def test_ai_resume_parser_reads_docx_tables():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Candidate Profile")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Python automation project experience"
    document.save(buffer)

    assert "automation project" in parse_resume("resume.docx", buffer.getvalue())


def test_ai_resume_parser_rejects_unknown_or_empty_files():
    with pytest.raises(ResumeParseError, match="只支持"):
        parse_resume("resume.txt", b"A long enough but unsupported resume body")
    with pytest.raises(ResumeParseError, match="为空"):
        parse_resume("resume.pdf", b"")
